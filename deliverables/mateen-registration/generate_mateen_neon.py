#!/usr/bin/env python3
"""Build the nine-slide Arabic Mateen concept pitch.

python -m pip install "python-pptx>=1.0.2,<2" pillow
python generate_mateen_neon.py --output mateen-neon-pitch.pptx

Uses the user-supplied official logo and a generated background only.
Install Cairo on the exporting machine; fonts are not embedded in PPTX.
All slide text, cards and diagrams are editable PowerPoint elements.
"""
import argparse
import math
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
ORIGINAL_LOGO = HERE.parents[1] / "attached_assets/MateeeeeeeeenLOGO_1790704101126.png"
LOGO = ORIGINAL_LOGO if ORIGINAL_LOGO.exists() else HERE / "mateen-official-logo.png"
W, H = 13.333333, 7.5
C = dict(bg="0B0F19", panel="1E293B", border="334155", blue="38BDF8",
         green="34D399", red="F87171", white="F8FAFC", muted="B5C4D7",
         subdued="8193AA", warm="C99A80")


def rgb(name):
    return RGBColor.from_string(C.get(name, name))


def rect(s, x, y, w, h, fill="panel", border=None, round_=False):
    shape = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if border:
        shape.line.color.rgb = rgb(border)
        shape.line.width = Pt(.75)
    else:
        shape.line.fill.background()
    if round_:
        shape.adjustments[0] = .07
    return shape


def text(s, value, x, y, w, h, size=21, color="white",
         bold=False, font="Cairo", align=PP_ALIGN.RIGHT, rtl=True):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_top = frame.margin_bottom = 0
    frame.margin_left = frame.margin_right = Inches(.018)
    for index, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.alignment = align
        p._p.get_or_add_pPr().set("rtl", "1" if rtl else "0")
        p.space_after = Pt(4)
        p.line_spacing = 1.06
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        rpr = run._r.get_or_add_rPr()
        rpr.set("lang", "ar-SA" if rtl else "en-US")
        cs = OxmlElement("a:cs")
        cs.set("typeface", font)
        rpr.append(cs)
    return shape


def make_assets(out):
    out.mkdir(parents=True, exist_ok=True)
    original = Image.open(LOGO).convert("RGBA")
    bbox = original.getbbox()
    if bbox is None:
        raise ValueError("The supplied official logo is empty.")
    official = out / "mateen-official-logo.png"
    original.crop(bbox).save(official)

    # Non-branded gradient is a background texture only, not a generated logo.
    width, height = 1600, 900
    image = Image.new("RGB", (width, height))
    px = image.load()
    for y in range(height):
        for x in range(width):
            t = y / (height - 1)
            base = (11 + 4*t, 15 + 8*t, 25 + 17*t)
            glow = math.exp(-(((x-1230)/490)**2 + ((y-210)/390)**2)*1.6)
            px[x, y] = (int(base[0]+1*glow), int(base[1]+12*glow),
                        int(base[2]+22*glow))
    bg = out / "neon-background.png"
    image.save(bg, optimize=True)
    return official, bg


def slide(prs, n, title, bg, logo, *, source=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(bg), 0, 0, width=Inches(W), height=Inches(H))
    # The logo itself is untouched; only transparent margins are cropped.
    s.shapes.add_picture(str(logo), Inches(11.66), Inches(.26),
                         width=Inches(.95))
    text(s, "المسار المفتوح  /  تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي",
         .65, .31, 9.8, .28, 9, "muted")
    rect(s, .65, .88, 12.02, .012, "border")
    text(s, title, .7, 1.12, 11.95, .76, 29, "white", True)
    rect(s, 11.76, 1.96, .86, .035, "blue")
    text(s, source or "مَتِين • عرض فكرة مقترحة للتسجيل، وليس عرض منتج منفّذ",
         1.5, 7.08, 11.1, .25, 8.5, "subdued")
    text(s, f"{n:02d} / 09", .65, 7.06, .73, .27, 10, "muted",
         rtl=False, align=PP_ALIGN.LEFT)
    return s


def side_label(s, value, x, y, color="blue"):
    text(s, value, x, y, 2.7, .35, 12, color, True)


def card(s, x, y, w, h, title, body, *, number=None, accent="blue"):
    rect(s, x, y, w, h, "panel", "border", True)
    rect(s, x+w-.06, y+.22, .025, h-.44, accent)
    if number:
        text(s, number, x+.23, y+.27, w-.47, .81, 42, accent, True)
        start = y+1.24
    else:
        start = y+.37
    text(s, title, x+.23, start, w-.48, .47, 19, "white", True)
    text(s, body, x+.23, start+.7, w-.48, h-(start-y)-.84,
         16.5, "muted")


def notes(s, value):
    s.notes_slide.notes_text_frame.text = value


def build(output):
    logo, bg = make_assets(output.parent)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "مَتِين | فكرة المسار المفتوح"
    prs.core_properties.subject = "تصور مبدئي للتحدي، لا بيان جاهزية تقنية"
    prs.core_properties.author = "فريق مَتِين"

    # 1 — 25s: full-bleed hero, original logo centered.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(bg), 0, 0, width=Inches(W), height=Inches(H))
    rect(s, 3.16, 1.15, 7.02, 1.96, "panel", "border", True)
    s.shapes.add_picture(str(logo), Inches(4.36), Inches(1.42),
                         width=Inches(4.62))
    text(s, "منصة مَتِين", 2, 3.35, 9.35, .76, 35, "white", True,
         align=PP_ALIGN.CENTER)
    text(s, "حفظ المتون والمنظومات العلمية، وفهمها بمصدر موثوق وتوجيه بشري",
         1.17, 4.24, 11, .67, 23, "muted", align=PP_ALIGN.CENTER)
    for x, label in [(8.43,"تسميع الكلمات"), (4.98,"شرح مسند"),
                     (1.52,"توجيه بشري")]:
        rect(s, x, 5.53, 3.02, .71, "panel", "border", True)
        text(s, label, x+.23, 5.68, 2.56, .38, 16, "blue", True,
             align=PP_ALIGN.CENTER)
    text(s, "تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي ٢٠٢٦ • المسار المفتوح",
         1.3, 7.02, 10.75, .31, 11, "muted", align=PP_ALIGN.CENTER)
    notes(s, "٢٥ ثانية. هذه فكرة وخطة تسجيل. مسارنا الرئيس المفتوح يجمع التعلم التفاعلي "
          "والحوار المعرفي الموثق؛ لا نسجل في المسارين الأول والثالث بالتزامن. "
          "الشعار الرسمي مرفق دون إعادة رسم.")

    # 2 — 40s: verified needs, not fabricated 61-respondent percentages.
    s = slide(prs, 2, "ثلاث فجوات في رحلة طالب المتون", bg, logo,
              source="استبيان الفريق: ٦٠ استجابة، ١٨–٢٢ سبتمبر ٢٠٢٦؛ الأسئلة ٢ و٣. عينة استطلاعية غير ممثلة إحصائياً.")
    text(s, "الاستبيان يحدد الحاجة، ولا يثبت أثر حل لم يُختبر بعد.",
         .74, 2.04, 11.6, .47, 17, "muted")
    card(s, 8.65, 2.85, 3.98, 3.23, "ضعف التفاعل",
         "٤٢ مشاركاً اختاروا الملل وفقدان الحافز ضمن عوائق الحفظ.",
         number="٧٠٪", accent="red")
    card(s, 4.67, 2.85, 3.76, 3.23, "صعوبة إيجاد شيخ",
         "٢٨ مشاركاً اختاروا صعوبة إيجاد من يسمّع لهم ويصحح.",
         number="٤٦٫٧٪")
    card(s, .7, 2.85, 3.75, 3.23, "تشتت الشروح",
         "٣٠ مشاركاً اختاروا طول الشروحات وتشتتها.",
         number="٥٠٪", accent="green")
    text(s, "نحتاج مساراً يصل الحفظ بالفهم والتوجيه، لا خدمة منفصلة لكل خطوة.",
         .75, 6.42, 11.65, .47, 19, "white", True)
    notes(s, "٤٠ ثانية. المقام ستون، لا واحد وستون. الأسئلة متعددة الاختيار؛ لا تجمع النسب. "
          "بعض المستجيبين لا يدرسون المتون حالياً، لذلك لا نزعم أنهم كلهم طلاب كليات شرعية. "
          "لا يوجد اقتباس حرفي موحد في البيانات كالعبارة المقترحة في الملف، لذلك لم ننسب لهم اقتباساً.")

    # 3 — 40s: scenario and clearly conceptual mini-interface.
    s = slide(prs, 3, "مَتِين: من حفظ النص إلى فهمه في رحلة واحدة", bg, logo)
    for i, (x, label, detail) in enumerate([
        (8.7, "١ / التسميع", "مطابقة الكلمات بالمتن؛ رصد الحذف والاستبدال بعد ثبوتهما."),
        (4.7, "٢ / الفهم", "شرح من شروح معتمدة مع إحالة دقيقة إلى المرجع."),
        (.7, "٣ / التوجيه", "إحالة نصية لمعلم معتمد أو سؤال مباشر عند الحاجة."),
    ]):
        card(s, x, 2.56, 3.74, 2.62, label, detail,
             accent=("blue", "green", "blue")[i])
    rect(s, .7, 5.55, 11.72, 1.23, "panel", "border", True)
    text(s, "تصور للتغذية الراجعة:  الكلمة الصحيحة  /  كلمة تحتاج إعادة  /  شرح مرتبط بالموضع",
         .99, 5.83, 11.12, .48, 19, "white")
    text(s, "النتيجة غير الواضحة تُطلب إعادتها؛ لا تُحسب خطأ حفظ مؤكداً.",
         1.02, 6.32, 11.04, .32, 15, "green")
    notes(s, "٤٠ ثانية. هذا مخطط واجهة تصوري وليس لقطة من تطبيق يعمل. النطاق الأول يقيّم "
          "صحة الكلمات وترتيبها والحذف والاستبدال. تقييم التشكيل والنطق ومخارج الحروف مؤجل. "
          "التصحيح الفوري في التدريب عند ثبوت الخطأ، دون إيقاف الطالب؛ لا يكشف الجواب في الاختبار.")

    # 4 — 40s: explicit 4-step workflow, no unsupported alignment model.
    s = slide(prs, 4, "كيف نقترح أن يعمل التسميع؟", bg, logo)
    steps = [
        (9.72, "٠١", "اختيار النص", "متن معتمد ومقطع محدد."),
        (6.73, "٠٢", "إرسال الصوت", "دفعات قصيرة مع حالة اتصال."),
        (3.74, "٠٣", "تفريغ ومطابقة", "مقارنة الكلمات بنص المقطع."),
        (.75, "٠٤", "نتيجة مفهومة", "تأكيد الخطأ أو طلب الإعادة."),
    ]
    for x, number, title, desc in steps:
        rect(s, x, 2.83, 2.84, 2.41, "panel", "border", True)
        text(s, number, x+.19, 3.04, 2.46, .61, 30, "blue", True)
        text(s, title, x+.19, 3.79, 2.46, .44, 17, "white", True)
        text(s, desc, x+.19, 4.41, 2.46, .66, 16, "muted")
        if x > 1:
            text(s, "←", x-.29, 3.67, .24, .44, 16, "blue")
    rect(s, .74, 5.65, 11.75, 1.03, "panel", "border", True)
    text(s, "مبدآن للتحقق: لا يعد الصمت وحده حذفاً؛ ولا يتحول صوت غير واضح إلى درجة أو خطأ نهائي.",
         1.01, 5.95, 11.18, .54, 19, "green", True)
    notes(s, "٤٠ ثانية. يبدأ المقترح بنص معتمد، وإذن الميكروفون، ودفعات مرقمة، وتفريغ "
          "وتطابق كلمات داخل المنصة. مقارنة مزودي الصوت على حذف واستبدال متعمدين وأخطاء زائفة "
          "وزمن ثبات النتيجة والكلفة. لا ادعاء بزمن أقل من ٣٠٠ مللي ثانية أو بنموذج forced alignment.")

    # 5 — 40s: safety is the product's differentiator, not an infallibility claim.
    s = slide(prs, 5, "الثقة العلمية تبدأ بحدود واضحة", bg, logo,
              source="دليل المشارك: الموثوقية والسالمة العلمية ٢٠٪ من معايير القبول؛ الاستبيان س٦.")
    card(s, 6.75, 2.62, 5.68, 3.49, "المساعد المسند",
         "بحث في مصادر معتمدة ومرخّصة.\nشرح محدود مع الكتاب والجزء والصفحة.\nالامتناع عند غياب المرجع أو تعارضه.",
         accent="blue")
    card(s, .7, 2.62, 5.68, 3.49, "الإحالة إلى المختص",
         "لا فتوى مولّدة من عند النظام.\nالسؤال المباشر متاح دون المرور بالمساعد.\nالسياق المشارك مع المعلم بموافقة الطالب.",
         accent="green")
    text(s, "الاسترجاع المقيد يخفف المخاطر ولا يجعل الإجابة معصومة من الخطأ.",
         .75, 6.42, 11.72, .48, 19, "red", True)
    notes(s, "٤٠ ثانية. استبيان سؤال ٦: ٤٧ من ٦٠، أي ٧٨٫٣٪ اختاروا ذكر المصدر "
          "ضمن عوامل الثقة. تؤخذ المصادر من حزمة علمية مأذون بها، وتراجع الطبعة وحقوق الاستخدام "
          "والمواضع المطبوعة قبل النشر. الحالات الحرجة: غياب المصدر، التعارض، وخارج النطاق والفتوى. "
          "لا ادعاء صفر هلوسة أو اكتمال مراجعة محتوى لم يُعتمد.")

    # 6 — 35s: architecture is prospective and close to approved specification.
    s = slide(prs, 6, "معمارية مقترحة، قابلة للاختبار والاستبدال", bg, logo)
    heads = [("الطبقة","الأداة المقترحة","وظيفتها")]
    rows = [
        ("الواجهة", "React + TypeScript + Tailwind", "قراءة عربية وتفاعل RTL"),
        ("الخادم والبيانات", "Node.js + Express + PostgreSQL / Drizzle", "صلاحيات وتقدم ومصادر"),
        ("الصوت", "مزود تفريغ قابل للتقييم + مطابقة كلمات", "نتيجة تدريجية ثم حكم مستقر"),
        ("المعرفة", "بحث نصي ودلالي + استرجاع مسند", "شرح بإحالة قابلة للتحقق"),
    ]
    rect(s, .75, 2.46, 11.78, .55, "blue", None, True)
    text(s, "الدور", 9.84, 2.53, 2.45, .36, 15, "bg", True)
    text(s, "الخيار", 4.78, 2.53, 4.8, .36, 15, "bg", True)
    text(s, "النتيجة", 1, 2.53, 3.54, .36, 15, "bg", True)
    for i, (role, tech, outcome) in enumerate(rows):
        yy = 3.12+i*.69
        rect(s, .75, yy, 11.78, .6, "panel", "border", True)
        text(s, role, 9.78, yy+.08, 2.48, .4, 14.5, "white", True)
        text(s, tech, 4.59, yy+.09, 4.98, .4, 13.5, "blue")
        text(s, outcome, 1.02, yy+.08, 3.34, .4, 14.5, "muted")
    text(s, "اختيار مزود الصوت والبحث النهائي يتبع اختبار الدقة والتكلفة وشروط الاحتفاظ، لا أسماء النماذج وحدها.",
         .82, 6.25, 11.5, .62, 18, "green")
    notes(s, "٣٥ ثانية. توافق اختيارات المواصفات المعتمدة. WebSocket للبث وREST للعمليات "
          "المعتادة، تخزين ملفات خاص لا قاعدة البيانات، وpgvector فقط عند دعم البيئة. "
          "تدريب نموذج متخصص يُدرس إذا أثبت التقييم ضرورة ذلك، لا شرط لرحلة أولى قابلة للتسليم.")

    # 7 — 40s: verified different denominators and meanings.
    s = slide(prs, 7, "ما قاله المستجيبون عن الثقة وتقبّل الفكرة", bg, logo,
              source="استبيان الفريق، ٦٠ استجابة، س٤–٦؛ الآراء عن فكرة مقترحة، لا رضا عن تجربة منتج.")
    card(s, 8.64, 2.81, 3.99, 3.37, "المصدر ضمن عوامل الثقة",
         "٤٧ من ٦٠ اختاروا ذكر المصدر الأصلي بوضوح.",
         number="٧٨٫٣٪", accent="blue")
    card(s, 4.67, 2.81, 3.75, 3.37, "المساعد والإحالة: ممتازة",
         "٣٨ من ٦٠ وصفوا هذا النموذج بأنه خيار ممتاز.",
         number="٦٣٫٣٪", accent="green")
    card(s, .7, 2.81, 3.75, 3.37, "التسميع الصوتي المقترح",
         "٣٥ من ٦٠ قيّموا فائدته المتوقعة ٤ أو ٥ من ٥.",
         number="٥٨٫٣٪", accent="blue")
    text(s, "إذا احتُسب «جيدة» أيضاً: ٥٤ من ٦٠ (٩٠٪) تقبّلوا فكرة المساعد مع الإحالة.",
         .74, 6.46, 11.7, .43, 17, "white")
    notes(s, "٤٠ ثانية. الحسابات من الملف الفعلي: ٤٧/٦٠=٧٨٫٣٪؛ ٣٨/٦٠=٦٣٫٣٪ "
          "لخيار ممتاز، لا ٦٣٫٩؛ ٣٥/٦٠=٥٨٫٣٪ لتقييم ٤ أو ٥، لا ٥٧٫٤. "
          "سؤال الصوت طلب تقييم ميزة تشمل تصحيح النطق والتشكيل؛ هذا أوسع من النطاق الأول "
          "الذي يقيم الكلمات، فلا تُستعمل النسبة لإثبات تبني ميزة الكلمات تحديداً. "
          "لم نستخدم شارات Code Ready أو Production Ready لأنها غير مثبتة.")

    # 8 — 25s: honest staged roadmap, not an already-launched V1.
    s = slide(prs, 8, "نبدأ ضيّقاً، ثم نتوسع بنتائج الاختبار", bg, logo)
    phases = [
        (8.65, "خلال التحدي", "متن مرخّص واحد؛ تسميع كلمات؛ شرح مسند؛ إحالة نصية."),
        (4.67, "بعد التحقق", "توسعة المتون المعتمدة واختبارات المستوى والمراجعة المتباعدة."),
        (.7, "لاحقاً", "بحث تقييم النطق والتشكيل والشراكات بعد تحقق البيانات والموارد."),
    ]
    for i, (x, title, detail) in enumerate(phases):
        rect(s, x, 2.79, 3.74, 3.04, "panel", "border", True)
        rect(s, x+.23, 3.15, .59, .59, "blue" if i == 0 else "border", None, True)
        text(s, f"٠{i+1}", x+.36, 3.25, .35, .32, 16, "bg" if i==0 else "white", True)
        text(s, title, x+.23, 3.94, 3.27, .47, 20, "white", True)
        text(s, detail, x+.23, 4.56, 3.27, 1.08, 17, "muted")
    rect(s, .78, 6.18, 11.7, .019, "blue")
    text(s, "الاستدامة: قياس تكلفة الصوت والسؤال والاستضافة، مع مسؤولية واضحة لمراجعة المصادر.",
         .83, 6.43, 11.53, .42, 18, "green")
    notes(s, "٢٥ ثانية. لا يوجد إطلاق حالي أو شراكات مبرمة. أثناء التحدي ٤–٦ أكتوبر "
          "نخطط لمتن واحد وشرح مرخّص بعد مراجعة الحقوق؛ بقية المواصفات المعتمدة على مراحل. "
          "لا تعد الشرائح بإصدار تلقائي للشهادات ولا بإطلاق مسار فقه أو لغة غير معتمد. "
          "تقاس كلفة دقيقة الصوت والسؤال والمراجعة البشرية بعد التجربة.")

    # 9 — 15s: official centered logo, actual team, no invented QR URLs.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(bg), 0, 0, width=Inches(W), height=Inches(H))
    s.shapes.add_picture(str(logo), Inches(4.83), Inches(.57),
                         width=Inches(3.66))
    text(s, "مَتِين: التقنية إلى جانب المصدر والمعلّم",
         1, 2.52, 11.25, .8, 30, "white", True, align=PP_ALIGN.CENTER)
    text(s, "نقترح اختبار رحلة تعلم متكاملة، بمؤشرات قابلة للتحقق وحدود معلنة.",
         1.2, 3.49, 10.87, .66, 21, "muted", align=PP_ALIGN.CENTER)
    rect(s, 6.8, 4.56, 5.55, 1.1, "panel", "border", True)
    text(s, "عبد الرحمن عماد ابوز عنين", 7.04, 4.73, 5.08, .35, 17, "blue", True)
    text(s, "الفكرة، التجربة، الواجهات والمصادر", 7.04, 5.12, 5.08, .31, 13, "muted")
    rect(s, .96, 4.56, 5.55, 1.1, "panel", "border", True)
    text(s, "مالك نورالدين الساعدي", 1.2, 4.73, 5.08, .35, 17, "green", True)
    text(s, "الذكاء الاصطناعي وتدريب النماذج والخادم", 1.2, 5.12, 5.08, .31, 13, "muted")
    text(s, "المخرجات المستهدفة عند نهاية التحدي: حل يعمل، ورابط تجريبي، ومستودع عام، وفيديو لا يتجاوز دقيقتين.",
         1.25, 6.15, 10.82, .76, 17, "white", align=PP_ALIGN.CENTER)
    notes(s, "١٥ ثانية. مجموع وقت الحديث المقترح ٣٠٠ ثانية. لا QR وهمياً ولا روابط مخترعة "
          "ولا أرقام اتصالات خاصة. الروابط يمكن إضافتها عند توفر رابط حي فعلي ومراجعة الحقوق "
          "حسب دليل التحدي. المسؤوليات وفق ما زودنا به الفريق دون اختراع شهادات.")

    assert len(prs.slides) == 9
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"Generated {output}: 9 slides, 16:9, 5-minute speaker notes.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path,
                        default=HERE / "mateen-neon-pitch.pptx")
    build(parser.parse_args().output)