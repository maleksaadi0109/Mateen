#!/usr/bin/env python3
"""Mateen concept pitch: exactly ten editable 16:9 Arabic slides.

pip install "python-pptx>=1.0.2,<2"
python generate_mateen_pitch.py --output mateen-concept-pitch.pptx

Install Kufam, Cairo and Noto Naskh Arabic for faithful rendering.
An optional --logo accepts a transparent cropped official logo.
No app, online API, survey file or other script is required.
"""
from pathlib import Path
import argparse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

PALETTE = dict(paper="FDF9F3", brown="6D4C3D", orange="994703",
               olive="4E3A00", ink="2A1F1A", card="F4EDE2")
FONTS = dict(title="Kufam", label="Cairo", body="Noto Naskh Arabic")
WIDTH, HEIGHT = 13.333333, 7.5


def rgb(value):
    return RGBColor.from_string(PALETTE.get(value, value))


def panel(s, x, y, w, h, fill="card", rounded=False, edge=None):
    a = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = rgb(fill)
    if edge:
        a.line.color.rgb = rgb(edge)
        a.line.width = Pt(.6)
    else:
        a.line.fill.background()
    if rounded:
        a.adjustments[0] = .08
    return a


def txt(s, value, x, y, w, h, size=22, ink="ink", font="body",
        bold=False, rtl=True, align=None):
    a = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t = a.text_frame
    t.clear()
    t.word_wrap = True
    t.margin_left = t.margin_right = Inches(.015)
    t.margin_top = t.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = t.paragraphs[0] if not i else t.add_paragraph()
        p.alignment = align if align is not None else PP_ALIGN.RIGHT
        p.space_after = Pt(6)
        p.line_spacing = 1.1
        p._p.get_or_add_pPr().set("rtl", "1" if rtl else "0")
        run = p.add_run()
        run.text = line
        run.font.name = FONTS[font]
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(ink)
        rp = run._r.get_or_add_rPr()
        rp.set("lang", "ar-SA" if rtl else "en-US")
        for tag in ("a:cs", "a:ea"):
            e = OxmlElement(tag)
            e.set("typeface", FONTS[font])
            rp.append(e)
    return a


def new(prs, n, section, dark=False, source=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb("brown" if dark else "paper")
    ink = "paper" if dark else "brown"
    txt(s, "منصة مَتِين  /  المسار المفتوح", 7, .28, 5.6, .32,
        11, ink, "label")
    txt(s, "تحدي المحتوى الإسلامي ٢٠٢٦  •  مقترح فكرة", .7, .28, 5.8, .32,
        10, ink, "label", align=PP_ALIGN.LEFT)
    panel(s, .7, .85, 11.9, .012, "card" if dark else "brown")
    txt(s, section, .7, 1.03, 11.9, .35, 12,
        "card" if dark else "orange", "label", True)
    txt(s, source or "عرض الفكرة للتسجيل • الوظائف وخطة التنفيذ مقترحة",
        1.4, 7.08, 11.2, .26, 9, ink, "label")
    txt(s, f"{n:02d}", .7, 7.03, .4, .34, 12, ink, "label", rtl=False)
    return s


def heading(s, text, dark=False, size=30):
    txt(s, text, .7, 1.65, 11.9, .95, size,
        "paper" if dark else "brown", "title", True)


def note(s, value):
    s.notes_slide.notes_text_frame.text = value


def tag(s, text, x, y, w, dark=False):
    panel(s, x, y, w, .43, "paper" if dark else "card", True)
    txt(s, text, x+.12, y+.05, w-.24, .31, 11, "orange", "label", True)


def build(output, logo=None):
    p = Presentation()
    p.slide_width, p.slide_height = Inches(WIDTH), Inches(HEIGHT)
    p.core_properties.title = "مَتِين — من الحفظ المنفرد إلى التعلم الموجّه"
    p.core_properties.subject = "عرض فكرة للتسجيل في المسار المفتوح"
    p.core_properties.author = "فريق مَتِين"

    # 1 — Editorial cover: proposition, not a feature inventory.
    s = new(p, 1, "حفظ المتون والمنظومات الشرعية • فهمها • مراجعتها")
    panel(s, 0, 1.55, 4.7, 5.2, "brown")
    for i, (a, b) in enumerate([
        ("أحفظ", "أعرف أين أخطأت."),
        ("أفهم", "أعود إلى مصدر واضح."),
        ("أسأل", "أصل إلى التوجيه البشري."),
    ]):
        yy = 2.04 + i*1.36
        txt(s, a, .68, yy, 3.3, .65, 31, "paper", "title", True)
        txt(s, b, .68, yy+.7, 3.3, .43, 21, "card")
    if logo and logo.exists():
        s.shapes.add_picture(str(logo), Inches(10.05), Inches(1.77), width=Inches(2.45))
    else:
        txt(s, "مَتِين", 7.2, 1.7, 5.2, 1.1, 52, "brown", "title", True)
    txt(s, "ليحفظ الطالب بثبات،\nويفهم بدليل.", 5.25, 3.05, 7.15, 1.8,
        38, "brown", "title", True)
    txt(s, "نقترح رحلة تجمع التسميع الذكي، والشرح المسند،\nوالتوجيه البشري؛ بدلاً من ترك الطالب بينها منفرداً.",
        5.35, 5.23, 7.05, 1.12, 24)
    note(s, "افتتاح مقترح: مَتِين ليست فكرة لإضافة كتاب آخر؛ بل لربط ما يحتاجه الطالب أثناء الدراسة: "
         "التسميع، وفهم العبارة، والوصول إلى المختص. نقدم تصوراً وخطة اختبار، لا نتائج تشغيل. "
         "لا ندعي الأسبقية السوقية أو ضمان تصحيح النطق والتشكيل.")

    # 2 — Evidence led, with one dominant finding and two supporting findings.
    s = new(p, 2, "الاحتياج الميداني • ٦٠ استجابة",
            source="استبيان الفريق، ١٨–٢٢ سبتمبر ٢٠٢٦، س٢–٣؛ المقام ٦٠. عينة استطلاعية غير ممثلة إحصائياً.")
    heading(s, "العائق ليس الحفظ وحده؛ بل الدراسة بلا تفاعل.")
    panel(s, .7, 2.87, 4.15, 3.87, "brown", True)
    txt(s, "٧٠٪", 1.07, 3.17, 3.42, 1.25, 70, "paper", "label", True)
    txt(s, "الملل وفقدان الحافز", 1.07, 4.7, 3.42, .5, 23, "paper", "label", True)
    txt(s, "٤٢ من ٦٠ اختاروا غياب التفاعل\nضمن عوائق الحفظ.", 1.07, 5.53, 3.42, .9,
        22, "paper")
    for yy, number, title, body in [
        (3.04, "٤٦٫٧٪", "شيخ متفرغ للتسميع", "٢٨ من ٦٠ اختاروا صعوبة إيجاده."),
        (4.95, "٥٠٪", "شروح طويلة ومشتتة", "٣٠ من ٦٠ اختاروا هذا العائق للفهم."),
    ]:
        txt(s, number, 10.15, yy, 2.42, .84, 34, "orange", "label", True)
        txt(s, title, 5.3, yy+.09, 4.45, .5, 23, "brown", "label", True)
        txt(s, body, 5.3, yy+.89, 7.24, .65, 22)
        panel(s, 5.3, yy+1.53, 7.25, .015, "card")
    note(s, "البيانات لا تقول إن كل طالب يعاني من كل مشكلة. الأسئلة متعددة الاختيارات؛ لا تجمع النسب. "
         "47 مشاركاً يدرسون متوناً أو منظومات و13 لا يدرسونها حالياً. خمس استجابات للسؤال الثاني "
         "تجاوزت اختيار عائقين؛ احتسبت كما وردت دون حذف صامت. لم تُنشر بيانات تعريفية.")

    # 3 — A conceptual learning journey, never an app screenshot.
    s = new(p, 3, "الحل • جمهور أول: دارسو المتون التأسيسية")
    heading(s, "كل موضع تعثّر يقابله إجراء واضح.")
    steps = [
        (8.72, "٠١", "عند التسميع", "هل حفظت الكلمات صحيحة؟",
         "مقارنة بالنص المعتمد لرصد\nالاستبدال والحذف بعد ثبوتهما."),
        (4.7, "٠٢", "عند الفهم", "ما معنى هذه العبارة؟",
         "شرح محدود بالمراجع،\nمع الكتاب والجزء والصفحة."),
        (.7, "٠٣", "عند حدود المعرفة", "مَن يوجّهني هنا؟",
         "إحالة نصية لمعلم معتمد،\nمع إتاحة السؤال المباشر."),
    ]
    for x, num, label, question, body in steps:
        panel(s, x, 2.97, 3.87, 3.23, "card", True)
        txt(s, num, x+.24, 3.15, 3.35, .66, 31, "orange", "label", True)
        txt(s, label, x+.24, 3.97, 3.35, .42, 16, "orange", "label")
        txt(s, question, x+.24, 4.52, 3.35, .46, 19, "brown", "label", True)
        txt(s, body, x+.24, 5.12, 3.35, 1.0, 19)
    txt(s, "واجهة عربية واضحة  •  تنبيه بالنص لا باللون وحده  •  حذف صوت التدريب بعد المعالجة",
        .8, 6.38, 11.7, .48, 18, "olive")
    note(s, "رحلة مفاهيمية وليست واجهة منتج. نبدأ بصحة الكلمات، لا تقييم مخارج الحروف أو الحركات. "
         "النتائج غير الواضحة تستدعي إعادة المقطع. الإحالة بموافقة الطالب وبالحد الأدنى من السياق؛ "
         "لا استنتاج لسمات دينية حساسة. مظهر فاتح وداكن، دون صور أو أيقونات لذوات الأرواح.")

    # 4 — Demand and scientific safeguards belong in the same argument.
    s = new(p, 4, "الثقة • شرط جوهري في الحل", True,
            "الاستبيان: س٦، ٤٧/٦٠ للمصدر؛ س٥، ٥٤/٦٠ للفكرة الجيدة أو الممتازة. تقبّل للفكرة لا اختبار للمنتج.")
    heading(s, "المصدر أولاً. وإذا لم يكفِ، فالإحالة.", True)
    for x, number, desc in [
        (7, "٧٨٫٣٪", "اختاروا ذكر المصدر ضمن عوامل الثقة."),
        (.8, "٩٠٪", "رأوا المساعد الموثق مع الإحالة فكرة جيدة أو ممتازة."),
    ]:
        txt(s, number, x, 2.92, 5.5, 1.28, 65, "paper", "label", True)
        txt(s, desc, x, 4.34, 5.5, .84, 23, "paper")
    panel(s, .8, 5.65, 11.75, 1.12, "paper", True)
    txt(s, "مرجع معتمد ← استشهاد قابل للتحقق ← إجابة محدودة أو امتناع",
        1.05, 5.84, 11.2, .46, 21, "brown", "label", True)
    txt(s, "لا فتوى آلية؛ ولا ادعاء بأن الاسترجاع المقيد يمنع الخطأ تماماً.",
        1.05, 6.36, 11.2, .3, 17, "orange")
    note(s, "خطة الموثوقية تمثل 20% من الترشيح. اختيار المحتوى من حزمة التحدي مع تدقيق الحقوق "
         "والطبعات والنص المستخرج؛ الكتاب والجزء والصفحة المطبوع لا رقم PDF بلا تمييز. "
         "اختبار المرجع المفقود والتعارض والسؤال خارج النطاق والفتوى. مراجعة علمية بشرية. "
         "الدليل يشير إلى مستويات محتوى أربعة دون تفصيلها في الملف؛ تُراجع الوثائق المنظمة عند إتاحتها.")

    # 5 — Two explainable pipelines instead of an acronym list.
    s = new(p, 5, "ملاءمة الذكاء الاصطناعي • قيمة محددة وليست إضافة شكلية")
    heading(s, "نحتاجه ليستمع ويسترجع؛ لا ليحكم بلا دليل.")
    lanes = [
        (2.97, "الصوت", ["صوت الطالب", "تفريغ عربي", "مطابقة الكلمات", "تنبيه أو إعادة"]),
        (4.6, "المعرفة", ["سؤال الطالب", "استرجاع مصدر", "تحقق من العزو", "شرح أو إحالة"]),
    ]
    for yy, label, cells in lanes:
        txt(s, label, 11.3, yy+.28, 1.23, .5, 20, "orange", "label", True)
        for i, value in enumerate(cells):
            xx = 8.64-i*2.64
            panel(s, xx, yy, 2.43, 1.08, "card", True)
            txt(s, value, xx+.13, yy+.32, 2.17, .52, 17, "brown", "label", True)
            if i<3:
                txt(s, "←", xx-.23, yy+.33, .23, .4, 15, "orange")
    txt(s, "React / TypeScript  •  Express / Node.js  •  PostgreSQL / Drizzle  •  WebSocket",
        .8, 6.05, 11.73, .34, 13, "orange", "label", rtl=False)
    txt(s, "معمارية مقترحة؛ اختيار النماذج بعد مقارنة الدقة والأخطاء الزائفة والتكلفة، لا التعقيد.",
        .8, 6.52, 11.73, .41, 18, "olive")
    note(s, "المقارنة بالكلمات منطق محدد لا قرار نموذج لغوي منفرد. نماذج الصوت خلف واجهة قابلة للاستبدال. "
         "الاسترجاع بحث نصي ودلالي وpgvector عند توافره. تدريب أو تكييف نموذج لا يُفترض ضرورياً "
         "خلال أيام التحدي؛ يقرر لاحقاً ببيانات مأذونة وقياسات تثبت الحاجة. فشل المزود لا يولد نجاحاً وهمياً.")

    # 6 — Honest comparison with a defined baseline.
    s = new(p, 6, "القيمة المضافة • لماذا المسار المفتوح؟")
    heading(s, "الإضافة ليست أداة أخرى؛ بل اتصال خطوات التعلم.")
    panel(s, .7, 2.85, 11.9, .61, "brown")
    txt(s, "مَتِين: ما نقترحه", .96, 2.98, 5.43, .35, 16, "paper", "label", True)
    txt(s, "الممارسة المنفصلة", 6.83, 2.98, 5.45, .35, 16, "paper", "label", True)
    comparisons = [
        ("حفظ منفرد دون تغذية راجعة أثناء المراجعة", "تسميع مرتبط بالنص ومواضع أخطاء الكلمات"),
        ("بحث يدوي بين الشروح لفهم عبارة", "شرح في سياق الدراسة مع مرجع يمكن الرجوع إليه"),
        ("سؤال المختص منفصل عن موضع التعثر", "إحالة تحمل العبارة والسؤال وسبب الحاجة إلى المختص"),
    ]
    for i, (old, proposed) in enumerate(comparisons):
        yy = 3.6+i*.78
        panel(s, .7, yy, 11.9, .7, "card" if i%2==0 else "paper")
        txt(s, proposed, .97, yy+.14, 5.44, .47, 20, "brown")
        txt(s, old, 6.83, yy+.14, 5.45, .47, 20)
    txt(s, "يجمع المقترح رحلة تعلم تفاعلية وحواراً معرفياً موثقاً لخدمة دارسي المحتوى الإسلامي.",
        .8, 6.15, 11.72, .64, 22, "orange")
    note(s, "ص12 من دليل المشارك يجيز للمفتوح جمع أكثر من مسار مع مشكلة وجمهور وأثر محدد. "
         "لا نضمن قبول التصنيف. نختار الممارسة الحالية معياراً للمقارنة دون الادعاء بتفوق على منتجات "
         "لم تختبر أو بعدم امتلاكها وظائف معينة. الفرضية: الربط يختصر الوصول إلى الفهم الصحيح؛ نختبرها ولا نفترضها.")

    # 7 — A realistic narrow delivery plan, not the whole product roadmap.
    s = new(p, 7, "قابلية التنفيذ • ٤–٦ أكتوبر ٢٠٢٦")
    heading(s, "لن نحاول بناء كل شيء في ثلاثة أيام.")
    panel(s, 9.12, 2.93, 3.48, 3.8, "brown", True)
    txt(s, "نطاق أول محدد", 9.4, 3.24, 2.92, .64, 23, "paper", "label", True)
    txt(s, "مقاطع من متن واحد\nشرح مرخّص ومراجع\nتسميع + شرح + إحالة",
        9.4, 4.16, 2.92, 1.62, 23, "paper")
    txt(s, "رحلة كاملة، لا وظائف صورية.", 9.4, 6.06, 2.92, .44, 16, "paper")
    days = [
        ("٤ أكتوبر", "المحتوى والصوت", "اعتماد المقاطع والمصدر؛ مقارنة مزودي التفريغ."),
        ("٥ أكتوبر", "تكامل الرحلة", "ربط التسميع والشرح والإحالة؛ اختبار حالات الفشل."),
        ("٦ أكتوبر", "التحقق والتسليم", "مراجعة علمية وتجربة متطوعين وتوثيق المخرجات."),
    ]
    for i, (date, title, body) in enumerate(days):
        yy = 3.03+i*1.22
        tag(s, date, 6.91, yy, 1.77)
        txt(s, title, .8, yy, 5.75, .45, 20, "brown", "label", True)
        txt(s, body, .8, yy+.57, 7.87, .53, 21)
    note(s, "خطة مقترحة مشروطة بالقبول وتوافر المصدر المرخص والمراجعة. مرشح من المحتوى المعتمد: "
         "الأربعون النووية؛ ليس اعتماداً لطبعة أو شرح قبل مراجعته. تؤجل بقية المحتويات والاختبارات الشاملة "
         "والتشكيل والنطق. بديل الصوت: مزود آخر أو معالجة التسجيل بعد انتهائه مع توضيح التأخير. "
         "عند عدم توافر معلم تعرض حالة الانتظار الحقيقية. مسؤولية المحتوى والتجربة لعبد الرحمن، "
         "والنماذج والخادم لمالك. لا تدريب من الصفر بوصفه شرطاً لهذه الأيام.")

    # 8 — Testable claims with explicit baselines and safety cases.
    s = new(p, 8, "معيار النجاح • ما الذي سنحاول إثباته؟")
    heading(s, "هل يفهم الطالب أسرع، دون أن نخسر الدقة؟")
    txt(s, "تجربة مقترحة: ١٠ متطوعين، مهام متكافئة، وترتيب متبادل بين مَتِين والبحث اليدوي.",
        .8, 2.74, 11.72, .65, 22)
    rows = [
        ("الفهم", "زمن الوصول إلى شرح صحيح ومسنَد", "مقارنة بالبحث اليدوي؛ لا يكفي أن تكون الإجابة أسرع."),
        ("التسميع", "اكتشاف الحذف والاستبدال", "٢٠ مقطعاً متنوعاً؛ قياس الأخطاء الزائفة وزمن التنبيه."),
        ("السلامة", "صحة العزو والامتناع والإحالة", "٢٠ سؤالاً؛ منها ١٠ حالات حرجة لمراجعة المختص."),
    ]
    for i, (label, metric, how) in enumerate(rows):
        yy = 3.61+i*.88
        panel(s, .7, yy, 11.9, .77, "card", True)
        txt(s, label, 10.79, yy+.2, 1.53, .4, 16, "orange", "label", True)
        txt(s, metric, 6.18, yy+.18, 4.27, .44, 19, "brown", "label", True)
        txt(s, how, .98, yy+.16, 4.83, .5, 18)
    txt(s, "هدف السلامة: اجتياز الحالات الحرجة قبل الإتاحة؛ لا أرقام أداء محققة ولا ضمان «صفر خطأ».",
        .8, 6.48, 11.72, .49, 18, "orange")
    note(s, "نطلب تجربة مضبوطة محدودة وليست إثبات أثر طويل الأجل. الحالات الحرجة: غياب المرجع، "
         "تعارضه، سؤال خارج النطاق، وطلب فتوى. يقيم المحتوى بمراجعة بشرية وفق إجابات مرجعية. "
         "يُقاس كذلك إتمام رحلة التسميع والشرح أو الإحالة. توثق النتائج والقيود والأخطاء، "
         "ولا تُستبدل بالمستهدفات الدعائية 95% أو 300ms.")

    # 9 — Real team, precise responsibilities, no invented credentials.
    s = new(p, 9, "الفريق • مسؤوليات واضحة من الفكرة إلى الاختبار")
    heading(s, "عضوان، وتكامل بين تجربة التعلم ومحركها.")
    for x, name, role, duties in [
        (6.82, "عبد الرحمن عماد ابوز عنين", "الفكرة • التجربة • المحتوى",
         "تطوير الفكرة وتجربة المستخدم\nتصميم الواجهات\nمراجعة صحة المحتوى وضبط المصادر"),
        (.7, "مالك نورالدين الساعدي", "الذكاء الاصطناعي • الخادم",
         "تطوير نماذج الذكاء الاصطناعي وتدريبها\nتطوير الواجهة الخلفية للمنصة\nربط مكونات المعالجة والخدمات"),
    ]:
        panel(s, x, 2.92, 5.77, 3.44, "card", True)
        panel(s, x+5.65, 3.18, .04, 2.86, "orange")
        txt(s, role, x+.27, 3.22, 5.17, .4, 13, "orange", "label", True)
        txt(s, name, x+.27, 3.93, 5.17, .7, 21, "brown", "title", True)
        txt(s, duties, x+.27, 4.94, 5.17, 1.24, 21)
    txt(s, "المراجعة المشتركة: جودة الرحلة، صحة الإسناد، وحدود النتائج قبل التسليم.",
        .8, 6.6, 11.72, .42, 19, "olive")
    note(s, "الأسماء والمسؤوليات كما قدمها الفريق، دون اختراع مؤهلات أو خبرات سابقة. "
         "تطوير النماذج وتدريبها مسؤولية مالك؛ لا يعني ذلك التزاماً بتدريب نموذج من الصفر في ثلاثة أيام. "
         "تُطلب الاستعانة بالإرشاد العلمي في التحدي، دون الادعاء بأن الفريق يضم مستشاراً شرعياً مستقلاً. "
         "المهمة المضافة «ربط المكونات» توضيح لنطاق تطوير الخادم لا اعتماد خبرة جديدة.")

    # 10 — A concrete closing proposition and a credible continuation plan.
    s = new(p, 10, "ما نريد اختباره في التحدي", True)
    txt(s, "حفظٌ لا ينفصل عن الفهم.\nوفهمٌ لا ينفصل عن المصدر.",
        .9, 1.92, 11.55, 1.82, 38, "paper", "title", True)
    txt(s, "نقترح مَتِين لاختبار رحلة تعلم واحدة متكاملة،\nمع نطاق صغير، ومسؤوليات واضحة، ونتائج يمكن التحقق منها.",
        .9, 4.04, 11.55, 1.0, 26, "paper")
    panel(s, .8, 5.47, 11.75, .67, "paper", True)
    txt(s, "المستهدف بنهاية التحدي: حل قابل للتجربة • مستودع عام • فيديو ≤ دقيقتين • عرض وسجل مصادر",
        1.03, 5.64, 11.27, .36, 15, "brown", "label", True)
    txt(s, "للاستمرار: قياس تكلفة دقيقة الصوت والسؤال، وضبط الاستخدام، ومراجعة المحتوى قبل توسيع المتون.",
        .9, 6.46, 11.55, .44, 19, "paper")
    note(s, "هذه مخرجات مستقبلية لنهاية التحدي وليست روابط أو مواد جاهزة الآن. "
         "خطة التشغيل تشمل تكلفة مزود الصوت والنموذج والتخزين والاستضافة والمراجعة العلمية؛ "
         "تقديراتها تُبنى على القياس، لا وعد بالمجانية. توافر خدمة WebSocket أثناء التحكيم وتجنب الخمول. "
         "توزيع الصيانة: مالك للخدمات والنماذج، وعبد الرحمن للمحتوى والتجربة. "
         "هذا العرض لا يستعرض أعمال تطوير ولا يقرر عدم وجودها؛ أي إفصاح تطلبه بوابة التسجيل "
         "عن نسخة البداية والحقوق يظل واجباً وفق ص43 من الدليل.")

    assert len(p.slides) == 10
    output.parent.mkdir(parents=True, exist_ok=True)
    p.save(output)
    print(f"Created {output}: 10 slides, 16:9")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path,
                        default=Path(__file__).with_name("mateen-concept-pitch.pptx"))
    parser.add_argument("--logo", type=Path)
    args = parser.parse_args()
    build(args.output, args.logo)