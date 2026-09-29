#!/usr/bin/env python3
"""Nine-slide editable Arabic Mateen pitch in the supplied brand identity.

Requires python-pptx and Pillow. The PPTX needs Kufam, Cairo and
Noto Naskh Arabic installed on the editing computer for exact typography.
"""
from pathlib import Path
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
W, H = 13.333333, 7.5
COLORS = {
    "cream": "FDF9F3", "paper": "F5F0E9", "card": "F3ECE3",
    "brown": "6D4C3D", "orange": "994703", "olive": "4E3A00",
    "ink": "2B2521", "muted": "685F57", "line": "DDD1C5",
    "white": "FFFFFF", "softbrown": "E9DBCD", "dark": "30251F",
    "red": "9A4B3E",
}
FONTS = {"head": "Kufam", "body": "Noto Naskh Arabic", "label": "Cairo"}


def c(name):
    return RGBColor.from_string(COLORS.get(name, name))


def box(slide, x, y, w, h, fill="card", border=None, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = c(fill)
    shape.line.color.rgb = c(border) if border else c(fill)
    shape.line.width = Pt(.85) if border else Pt(0)
    if rounded:
        shape.adjustments[0] = .06
    return shape


def txt(slide, value, x, y, w, h, size=20, color="ink", font="body",
        bold=False, align=PP_ALIGN.RIGHT, rtl=True, line_spacing=1.04):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = sh.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.015)
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p._p.get_or_add_pPr().set("rtl", "1" if rtl else "0")
        p.space_after = Pt(2)
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.name = FONTS[font]
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = c(color)
        rp = r._r.get_or_add_rPr()
        rp.set("lang", "ar-SA" if rtl else "en-US")
        for tag in ("a:cs", "a:ea"):
            el = OxmlElement(tag)
            el.set("typeface", FONTS[font])
            rp.append(el)
    return sh


def photo(slide, source, x, y, w, h):
    # Pre-cropped screenshot remains an image. No fabricated UI is inserted.
    slide.shapes.add_picture(str(source), Inches(x), Inches(y),
                             width=Inches(w), height=Inches(h))


def notes(slide, speech):
    slide.notes_slide.notes_text_frame.text = speech


def header(prs, number, title, logo, *, eyebrow="تصور المشروع • المسار المفتوح",
           foot="مَتِين  |  عرض مشروع • الوظائف والخطط مقترحة ما لم يُذكر خلاف ذلك"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = c("cream")
    box(s, 0, 0, W, .105, "brown", rounded=False)
    photo(s, logo, 11.62, .23, .93, .44)
    txt(s, eyebrow, .7, .29, 10.43, .28, 10, "brown", "label", True)
    box(s, .72, .86, 11.9, .014, "line", rounded=False)
    txt(s, title, .72, 1.02, 11.88, .9, 29, "brown", "head", True)
    box(s, 11.85, 1.96, .73, .045, "orange", rounded=False)
    txt(s, foot, 1.38, 7.11, 11.2, .23, 9, "muted", "label")
    txt(s, f"{number:02d} / 09", .73, 7.08, .68, .26, 10,
        "brown", "label", rtl=False, align=PP_ALIGN.LEFT)
    return s


def small_card(s, x, y, w, h, heading, content, accent="orange",
               font_size=19):
    box(s, x, y, w, h, "white", "line")
    box(s, x+w-.075, y+.23, .045, h-.46, accent, rounded=False)
    txt(s, heading, x+.24, y+.29, w-.56, .48, 17.5,
        accent, "label", True)
    txt(s, content, x+.24, y+.9, w-.57, h-1.05, font_size, "ink")


def build_assets():
    logo = HERE / "mateen-official-logo.png"
    if not logo.exists():
        source = HERE.parents[1] / "attached_assets/MateeeeeeeeenLOGO_1790704101126.png"
        img = Image.open(source).convert("RGBA")
        img.crop(img.getbbox()).save(logo)
    capture = HERE / "project-home.jpg"
    older = HERE.parents[1] / "attached_assets/image_1790092620427.png"
    if not capture.exists() or not older.exists():
        raise FileNotFoundError("Actual Mateen project screenshots are required.")
    im = Image.open(capture).convert("RGB")
    ImageOps.fit(im.crop((50, 155, 775, 885)), (1250, 780)).save(
        HERE / "project-recitation-crop.jpg", quality=91)
    im = Image.open(older).convert("RGB")
    ImageOps.fit(im.crop((780, 55, 1915, 960)), (1250, 780)).save(
        HERE / "project-assistant-crop.jpg", quality=91)
    return logo


def build():
    logo = build_assets()
    p = Presentation()
    p.slide_width, p.slide_height = Inches(W), Inches(H)
    p.core_properties.title = "مَتِين | عرض الفكرة بهوية المشروع"
    p.core_properties.subject = "عرض من تسع شرائح للمسار المفتوح"
    p.core_properties.author = "فريق مَتِين"

    # 1 — Preserve the centered hero structure the user liked; use brand colors.
    s = p.slides.add_slide(p.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = c("cream")
    box(s, 0, 0, W, .13, "brown", rounded=False)
    box(s, 2.95, .88, 7.43, 2.13, "paper", "line")
    # Keep the official logo's original 622:287 proportions.
    photo(s, logo, 5.02, 1.19, 3.30, 1.52)
    txt(s, "منصة مَتِين", 2, 3.27, 9.34, .72, 35,
        "brown", "head", True, PP_ALIGN.CENTER)
    txt(s, "حفظ المتون والمنظومات العلمية، وفهمها بمصدر موثوق وتوجيه بشري",
        1.2, 4.2, 10.94, .73, 24, "ink", align=PP_ALIGN.CENTER)
    for x, label in [(8.46, "تسميع الكلمات"), (5.02, "شرح مسند"),
                     (1.58, "توجيه بشري")]:
        box(s, x, 5.45, 2.93, .73, "white", "line")
        txt(s, label, x+.18, 5.63, 2.55, .38, 17,
            "orange", "label", True, PP_ALIGN.CENTER)
    txt(s, "تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي ٢٠٢٦  •  المسار المفتوح",
        1.2, 6.92, 10.96, .36, 12, "brown", "label",
        align=PP_ALIGN.CENTER)
    notes(s, "٢٥ ثانية. نعرض الفكرة وهوية مَتِين. التسميع، الفهم المسند، والتوجيه البشري "
          "مرتكزات رحلة التعلم المقترحة. نحن في المسار المفتوح. الصور اللاحقة من واجهة المشروع "
          "ولا تُعد بحد ذاتها إثباتاً لاختبار كل وظيفة.")

    # 2 — User-provided four-axis concept, with verified survey figures.
    s = header(p, 2, "مَتِين | بيئة تعلم متكاملة بمرجعية واضحة", logo)
    txt(s, "تُعد منصة «مَتِين» بيئة تعليمية تقنية تهدف إلى تيسير حفظ وتسميع وفهم المتون "
        "والمنظومات العلمية بدمج تقنيات الذكاء الاصطناعي مع الإشراف العلمي البشري المعتمد.",
        .77, 1.99, 11.8, .77, 19.5, "ink")
    txt(s, "تتجسد أصالة المنصة ومرجعيتها العلمية والتقنية في أربعة محاور:",
        .77, 2.77, 11.8, .42, 17, "orange", "label", True)
    small_card(s, 6.81, 3.23, 5.71, 1.65, "٠١  التوثيق الشرعي الصارم  |  Strict RAG",
               "إجابة من الشروح المحققة المعتمدة، مع اسم الكتاب والصفحة؛ "
               "امتناع عن الإفتاء والتكهن.", font_size=17.2)
    small_card(s, .75, 3.23, 5.71, 1.65, "٠٢  التسميع الصوتي الذكي",
               "تسميع مقطعاً بمقطع؛ مطابقة كلمات لرصد الزيادة والنقص والترتيب، "
               "لا حكم من نموذج محادثة عام.", font_size=17.2)
    small_card(s, 6.81, 5.04, 5.71, 1.58, "٠٣  الإشراف البشري  |  Human-in-the-Loop",
               "إحالة الحالات المعقدة والاستفسارات المتقدمة إلى "
               "معلم معتمد؛ والسؤال المباشر متاح.", font_size=17.2)
    small_card(s, .75, 5.04, 5.71, 1.58, "٠٤  الواقعية والتحقق الميداني",
               "استبيان من ٦٠ استجابة: ٧٨٫٣٪ اختاروا ذكر المصدر "
               "ضمن عوامل الثقة، و٦٣٫٣٪ وصفوا المساعد مع الإحالة بأنه ممتاز.",
               font_size=16.8)
    notes(s, "٤٠ ثانية. صُححت الأرقام من الملف المرفق فعلياً: ٦٠ استجابة، ٤٧ من ٦٠ = "
          "٧٨٫٣ بالمئة اختاروا ذكر المصدر ضمن عوامل الثقة، و٣٨ من ٦٠ = ٦٣٫٣ بالمئة "
          "وصفوا الفكرة بأنها ممتازة. لا نعرض هذه الآراء بوصفها نتائج منتج مُختبر. "
          "نقترح قصر المساعد على المصادر المعتمدة والامتناع والإحالة عند الحاجة؛ "
          "لا ندعي منع الخطأ بصورة مطلقة.")

    # 3 — The problem, with calibrated exploratory survey data.
    s = header(p, 3, "المشكلة | رحلة الطالب موزعة بين الحفظ والفهم والتوجيه", logo,
               foot="المصدر: استبيان الفريق، ٦٠ استجابة، س٢–٣؛ متعدد الاختيارات وعينته استطلاعية.")
    txt(s, "ثلاث فجوات مترابطة ظهرت في استجابات العينة:",
        .79, 2.07, 11.7, .43, 19, "muted", "label")
    for x, n, title, detail, color in [
        (8.6, "٧٠٪", "ضعف التفاعل", "٤٢ من ٦٠ اختاروا الملل وفقدان الحافز.", "brown"),
        (4.66, "٤٦٫٧٪", "صعوبة إيجاد معلّم", "٢٨ من ٦٠ ذكروا صعوبة إيجاد من يُسمّع ويصحح.", "orange"),
        (.72, "٥٠٪", "تشتت الشروح", "٣٠ من ٦٠ اختاروا طول الشروح وتشتتها.", "olive"),
    ]:
        box(s, x, 2.78, 3.77, 3.3, "white", "line")
        txt(s, n, x+.26, 3.1, 3.26, .85, 43, color, "label", True)
        txt(s, title, x+.26, 4.14, 3.27, .49, 20, "brown", "label", True)
        txt(s, detail, x+.26, 4.88, 3.24, .88, 19, "ink")
    box(s, .75, 6.31, 11.77, .52, "softbrown")
    txt(s, "المطلوب: وصل التسميع بالشرح المسند ثم التوجيه، بدل ترك كل خطوة منفصلة.",
        1.03, 6.41, 11.2, .33, 17, "brown", "label", True)
    notes(s, "٤٠ ثانية. هذه استطلاع آراء، لا قياس أثر، والأسئلة متعددة الاختيار فلا تُجمع النسب. "
          "لا ننسب لكل المستجيبين صفة طالب كلية شرعية. البيانات تشير إلى فجوات متصلة في التعلم.")

    # 4 — Real, user-supplied and current screenshots, clearly labeled.
    s = header(p, 4, "الحل | واجهة تجمع التسميع والمساعد في سياق واحد", logo,
               foot="لقطات من واجهة مشروع مَتِين؛ لا تُستخدم لإثبات دقة نموذج الصوت أو صحة جميع الإجابات.")
    txt(s, "صور المشروع", .78, 1.98, 11.7, .4, 17, "orange", "label", True)
    for x, image, label in [
        (6.79, HERE / "project-recitation-crop.jpg", "واجهة التسميع والنص الدراسي"),
        (.75, HERE / "project-assistant-crop.jpg", "واجهة المساعد والسؤال"),
    ]:
        box(s, x, 2.51, 5.73, 3.45, "white", "line")
        photo(s, image, x+.14, 2.65, 5.45, 2.76)
        txt(s, label, x+.2, 5.52, 5.28, .33, 17, "brown", "label", True)
    txt(s, "المبدأ: يراجع الطالب موضعه، ثم يطلب شرحاً مسنداً أو توجيهاً بشرياً عند الحاجة.",
        .82, 6.23, 11.56, .65, 19, "ink")
    notes(s, "٤٠ ثانية. الصورتان لواجهة المشروع: لقطة حديثة للتسميع ولقطة مقدمة سابقاً "
          "من واجهة المساعد. شكل الشاشة لا يثبت وحده اكتمال التقييم الصوتي أو المرجعية العلمية؛ "
          "هذه الأمور تُختبر وتوثق مستقلاً قبل تقديم نتائج تشغيل.")

    # 5 — Mechanism.
    s = header(p, 5, "آلية العمل | من صوت الطالب إلى قرار قابل للفهم", logo)
    steps = [
        (9.68, "٠١", "تحديد المقطع", "متن محقق ومقطع واضح للنطاق."),
        (6.70, "٠٢", "استقبال الصوت", "تسجيل مقاطع قصيرة بإذن الطالب."),
        (3.72, "٠٣", "تفريغ ومطابقة", "كشف حذف أو استبدال أو زيادة الكلمات."),
        (.74, "٠٤", "تنبيه أو إعادة", "نتيجة مؤكدة؛ وغير الواضح يُعاد."),
    ]
    for x, num, title, body in steps:
        box(s, x, 2.78, 2.84, 2.75, "white", "line")
        txt(s, num, x+.21, 3.01, 2.44, .61, 29, "orange", "label", True)
        txt(s, title, x+.21, 3.87, 2.42, .42, 17, "brown", "label", True)
        txt(s, body, x+.21, 4.5, 2.44, .87, 18, "ink")
        if x > 1:
            txt(s, "←", x-.28, 3.92, .26, .44, 18, "orange", "label")
    box(s, .75, 5.96, 11.72, .86, "softbrown")
    txt(s, "لا يُحتسب الصمت وحده حذفاً، ولا تتحول قراءة غير واضحة إلى خطأ حفظ مؤكد.",
        1, 6.15, 11.2, .46, 19, "brown", "label", True)
    notes(s, "٤٠ ثانية. هذا مسار عمل مقترح. يبدأ باختيار متن معتمد، وتقسيم الصوت وتفريغه "
          "ومقارنة الكلمات بنص المقطع. نقيم الحذف والزيادة والترتيب بعد ثبوتها؛ "
          "النطق والتشكيل ومخارج الحروف ليست حكماً أولياً. لا نعد بزمن تنبيه لم نقسه.")

    # 6 — Differentiation and safety.
    s = header(p, 6, "القيمة المضافة | كل خطوة تقود إلى التالية", logo)
    box(s, .75, 2.36, 11.78, .64, "brown")
    txt(s, "الممارسة المتفرقة", 7.05, 2.51, 5.14, .37, 17, "white", "label", True)
    txt(s, "ما تقترحه مَتِين", 1.14, 2.51, 5.42, .37, 17, "white", "label", True)
    for i, (old, new) in enumerate([
        ("حفظ فردي دون موضع واضح للخطأ", "تسميع مرتبط بالمقطع وموضع الكلمة"),
        ("بحث منفصل بين الشروح", "شرح في سياق النص مع عزو للكتاب والصفحة"),
        ("سؤال غير مرتبط بموضع التعثر", "إحالة تحمل سؤال الطالب وسياقه بموافقته"),
    ]):
        y = 3.13 + i*.83
        box(s, .75, y, 11.78, .72, "white", "line")
        txt(s, old, 6.82, y+.15, 5.28, .43, 18, "muted")
        txt(s, new, 1.02, y+.15, 5.36, .43, 18, "brown", bold=True)
    box(s, .75, 5.87, 11.78, .93, "paper", "line")
    txt(s, "السلامة العلمية: مصدر معتمد ← استشهاد قابل للتحقق ← إجابة محدودة أو امتناع وإحالة.",
        1.06, 6.08, 11.15, .5, 19, "olive", "label", True)
    notes(s, "٣٥ ثانية. لا ندعي أن منتجات أخرى لا تقدم هذه المزايا؛ المقارنة مع "
          "الخطوات المنفصلة لدى الطالب. الإسناد والتحقق والامتناع والإحالة متطلبات للسلامة "
          "وليست شهادة بعدم وجود أخطاء. لا فتوى مولدة من النظام.")

    # 7 — Technical architecture.
    s = header(p, 7, "التقنيات | معمارية قابلة للاختبار والاستبدال", logo)
    box(s, .76, 2.35, 11.78, .59, "brown")
    txt(s, "الطبقة", 9.69, 2.48, 2.58, .34, 16, "white", "label", True)
    txt(s, "الخيار المقترح", 4.41, 2.48, 5.1, .34, 16, "white", "label", True)
    txt(s, "الغاية", 1.03, 2.48, 3.17, .34, 16, "white", "label", True)
    for i, (layer, tech, goal) in enumerate([
        ("الواجهة", "React + TypeScript + Tailwind", "قراءة وتفاعل عربي RTL"),
        ("الخادم والبيانات", "Node.js + Express + PostgreSQL / Drizzle", "صلاحيات ومحتوى وتقدم"),
        ("الصوت", "تفريغ قابل للمقارنة + مطابقة كلمات", "رصد الخطأ بعد التأكد"),
        ("المعرفة", "بحث نصي ودلالي + استرجاع مسند", "شرح بعزو يمكن مراجعته"),
    ]):
        y = 3.11 + i*.75
        box(s, .76, y, 11.78, .66, "white", "line")
        txt(s, layer, 9.67, y+.13, 2.56, .4, 16, "brown", "label", True)
        txt(s, tech, 4.22, y+.15, 5.28, .38, 15, "orange", "label")
        txt(s, goal, 1.02, y+.14, 3.12, .39, 16.5, "ink")
    txt(s, "يُحسم اختيار مزود الصوت بعد قياس دقة الكشف، الأخطاء الزائفة، التكلفة والاحتفاظ بالبيانات.",
        .84, 6.35, 11.52, .51, 18, "olive")
    notes(s, "٣٥ ثانية. هذه اختيارات معمارية مقترحة وليست شهادة بأن كل مسار يعمل بإنتاجية. "
          "الصوت خلف واجهة تسمح بتبديل مزوده. البحث في المصادر المأذون بها والتحقق من الطبعات "
          "والصفحات جزء من الخطة. التدريب المتخصص يُبحث بعد اختبار الحاجة إليه.")

    # 8 — Field findings, not product outcomes.
    s = header(p, 8, "النتائج | تقبّل الفكرة لا يساوي نجاح المنتج بعد", logo,
               foot="المصدر: استبيان الفريق، ٦٠ استجابة، الأسئلة ٤–٦. النتائج تفضيلات معلنة وليست أثر استخدام.")
    for x, number, title, detail, accent in [
        (8.61, "٧٨٫٣٪", "ذكر المصدر", "٤٧ من ٦٠ اختاروا ذكر المرجع ضمن عوامل الثقة.", "brown"),
        (4.65, "٦٣٫٣٪", "المساعد والإحالة", "٣٨ من ٦٠ وصفوا هذا التصور بأنه ممتاز.", "orange"),
        (.71, "٥٨٫٣٪", "صوت مقترح", "٣٥ من ٦٠ قيّموا فكرة صوتية أوسع بـ٤–٥ من ٥.", "olive"),
    ]:
        box(s, x, 2.8, 3.79, 3.37, "white", "line")
        txt(s, number, x+.25, 3.16, 3.3, .93, 40, accent, "label", True)
        txt(s, title, x+.25, 4.32, 3.26, .45, 19, "brown", "label", True)
        txt(s, detail, x+.25, 5.05, 3.28, .91, 18, "ink")
    txt(s, "الخطوة التالية: اختبار دقة التسميع وصحة العزو وإتمام الرحلة قبل إعلان نتائج أداء.",
        .81, 6.42, 11.55, .52, 18, "brown", "label", True)
    notes(s, "٣٠ ثانية. ٤٧/٦٠ = ٧٨٫٣ بالمئة اختاروا المصدر ضمن عوامل الثقة؛ "
          "٣٨/٦٠ = ٦٣٫٣ وصفوا فكرة المساعد مع الإحالة بأنها ممتازة؛ ٣٥/٦٠ = ٥٨٫٣ "
          "قيّموا الصوت ٤ أو ٥، لكن السؤال يشمل النطق والتشكيل، وهو أوسع من المرحلة الأولى. "
          "هذا استطلاع احتياج وتقبل وليس نتيجة أداء أو دليل فعالية.")

    # 9 — Continuation and team / closing.
    s = header(p, 9, "الاستمرار | نثبت رحلة واحدة ثم نتوسع بمسؤولية", logo)
    for x, num, title, detail in [
        (8.64, "٠١", "نطاق أول", "متن معتمد واحد؛ كلمات وشرح مسند وإحالة نصية."),
        (4.66, "٠٢", "قياس وتطوير", "دقة الكشف والعزو؛ تكلفة الصوت والسؤال والمراجعة."),
        (.72, "٠٣", "توسع مشروط", "متون إضافية بعد الترخيص والاختبار والتدقيق."),
    ]:
        box(s, x, 2.39, 3.77, 2.38, "white", "line")
        txt(s, num, x+.22, 2.64, 3.3, .47, 22, "orange", "label", True)
        txt(s, title, x+.22, 3.21, 3.29, .41, 18, "brown", "label", True)
        txt(s, detail, x+.22, 3.84, 3.29, .8, 18, "ink")
    box(s, .74, 5.12, 11.8, 1.17, "brown")
    txt(s, "عبد الرحمن عماد ابوز عنين", 6.76, 5.35, 5.49, .39,
        17, "white", "label", True)
    txt(s, "الفكرة والتجربة والواجهات ومراجعة المصادر", 6.76, 5.83, 5.49, .28,
        13.5, "white", "label")
    txt(s, "مالك نورالدين الساعدي", 1.05, 5.35, 5.26, .39,
        17, "white", "label", True)
    txt(s, "الذكاء الاصطناعي والنماذج والخادم", 1.05, 5.83, 5.26, .28,
        13.5, "white", "label")
    txt(s, "مَتِين: تقنية مساعدة، ومصدر يمكن التحقق منه، ومعلّم لا يغيب عن القرار.",
        1.12, 6.54, 11.05, .37, 18, "olive", "label", True,
        align=PP_ALIGN.CENTER)
    notes(s, "١٥ ثانية. التنفيذ والتوسع مشروطان بالاعتماد والقياس. نراجع حقوق المحتوى "
          "والمصدر والدقة والتكلفة قبل التوسعة. فريق المشروع كما قدم الأسماء والأدوار. "
          "لا نعطي رابط تجربة أو مستودعاً افتراضياً.")

    assert len(p.slides) == 9
    dest = HERE / "mateen-brand-pitch.pptx"
    p.save(dest)
    print(f"Created {dest} (9 slides, editable elements, 5-minute notes)")


if __name__ == "__main__":
    build()